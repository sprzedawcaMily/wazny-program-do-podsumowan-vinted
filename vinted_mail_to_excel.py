import email
import imaplib
import json
import os
import re
import ssl
import subprocess
import time
import webbrowser
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from email.header import decode_header
from email.utils import parsedate_to_datetime
from html import unescape
from typing import Dict, List, Optional, Tuple

from deep_translator import GoogleTranslator
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


SALE_SUBJECT = "zamowienie zostalo zakonczone"
PURCHASE_SUBJECT = "twoje potwierdzenie zakupu"
SERVICE_SUBJECT = "faktura za usluge ekspozycja"


@dataclass
class Record:
    kind: str
    tx_id: str
    email_uid: str
    subject: str
    title_original: str
    title_pl: str
    amount_total: str
    date_text: str
    country: str
    shipping_international: str
    vat_number: str
    transaction_number: str


def normalize_text(value: str) -> str:
    replacements = {
        "ą": "a",
        "ć": "c",
        "ę": "e",
        "ł": "l",
        "ń": "n",
        "ó": "o",
        "ś": "s",
        "ż": "z",
        "ź": "z",
        "Ą": "A",
        "Ć": "C",
        "Ę": "E",
        "Ł": "L",
        "Ń": "N",
        "Ó": "O",
        "Ś": "S",
        "Ż": "Z",
        "Ź": "Z",
    }
    cleaned = re.sub(r"[<>:\"/\\|?*]", "_", value)
    text = cleaned
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    return text


def decode_mime_header(value: Optional[str]) -> str:
    if not value:
        return ""
    decoded_parts = decode_header(value)
    chunks: List[str] = []
    for part, enc in decoded_parts:
        if isinstance(part, bytes):
            chunks.append(part.decode(enc or "utf-8", errors="replace"))
        else:
            chunks.append(part)
    return "".join(chunks).strip()


def html_to_text(html: str) -> str:
    text = re.sub(r"(?is)<(script|style).*?>.*?(</\\1>)", " ", html)
    text = re.sub(r"(?is)<br\\s*/?>", "\n", text)
    text = re.sub(r"(?is)</p>", "\n", text)
    text = re.sub(r"(?is)<.*?>", " ", text)
    text = unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s+", "\n", text)
    text = re.sub(r"\n{2,}", "\n\n", text)
    return text.strip()


def extract_plain_text(msg: email.message.Message) -> str:
    if msg.is_multipart():
        plain_parts: List[str] = []
        html_parts: List[str] = []
        for part in msg.walk():
            ctype = part.get_content_type()
            disp = str(part.get("Content-Disposition", ""))
            if "attachment" in disp.lower():
                continue
            payload = part.get_payload(decode=True)
            if payload is None:
                continue
            charset = part.get_content_charset() or "utf-8"
            text = payload.decode(charset, errors="replace")
            if ctype == "text/plain":
                plain_parts.append(text)
            elif ctype == "text/html":
                html_parts.append(text)
        if plain_parts:
            return "\n".join(plain_parts).strip()
        if html_parts:
            return "\n".join(html_to_text(h) for h in html_parts).strip()
        return ""

    payload = msg.get_payload(decode=True)
    if payload is None:
        return ""
    charset = msg.get_content_charset() or "utf-8"
    raw = payload.decode(charset, errors="replace")
    if msg.get_content_type() == "text/html":
        return html_to_text(raw)
    return raw


def parse_date_from_email(msg: email.message.Message, body: str) -> str:
    date_hdr = msg.get("Date")
    if date_hdr:
        try:
            dt = parsedate_to_datetime(date_hdr)
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=timezone.utc)
            return dt.astimezone().strftime("%d.%m.%Y")
        except Exception:
            pass

    m = re.search(r"Data\s*:\s*(\d{1,2}\.\d{1,2}\.\d{4})", body, flags=re.IGNORECASE)
    if m:
        dd, mm, yyyy = m.group(1).split(".")
        return f"{dd.zfill(2)}.{mm.zfill(2)}.{yyyy}"

    return ""


def parse_country(body: str) -> str:
    text = body.replace("\r", "")
    match = re.search(r"Vinted\s*,\s*UAB(.{0,400})", text, flags=re.IGNORECASE | re.DOTALL)
    if not match:
        return ""

    block = match.group(0)
    block = re.split(r"sprzedawc|kupuj|otrzymano|przelano", block, flags=re.IGNORECASE)[0]
    lines = []
    for raw_line in block.splitlines():
        left = raw_line.split("|")[0].strip()
        if left:
            lines.append(left)

    if not lines:
        lines = [block.strip()]

    return _extract_country_from_block(lines)


def _extract_country_from_block(lines: List[str]) -> str:
    cleaned: List[str] = []
    for ln in lines:
        if not ln:
            continue
        normalized = normalize_text(ln).lower()
        if "vat" in normalized or re.search(r"\b[a-z]{2}\d{6,}\b", normalized):
            continue
        cleaned.append(ln)

    for ln in cleaned:
        normalized = normalize_text(ln).lower()
        if re.search(r"\b(litwa|lithuania)\b", normalized):
            return "Litwa"
        if re.search(r"\b(polska|poland)\b", normalized):
            return "Polska"

    for ln in cleaned:
        code = normalize_text(ln).strip().lower()
        if code in {"lt"}:
            return "Litwa"
        if code in {"pl"}:
            return "Polska"

    for ln in reversed(cleaned):
        if not ln:
            continue
        normalized = normalize_text(ln).lower()
        if re.search(r"\blitwa\b|\blithuania\b", normalized) or normalized.strip() in {"lt"}:
            return "Litwa"
        if re.search(r"\bpolska\b|\bpoland\b", normalized) or normalized.strip() in {"pl"}:
            return "Polska"

        m_code = re.search(r"\b(lt|pl)-\d{2,}\b", normalized)
        if m_code:
            return "Litwa" if m_code.group(1) == "lt" else "Polska"

        m = re.search(r",\s*([A-Za-z\-ĄĆĘŁŃÓŚŻŹąćęłńóśżź]+)\s*$", ln)
        if m:
            candidate = m.group(1).strip()
            candidate_norm = normalize_text(candidate).lower()
            if candidate_norm in {"pl", "polska", "poland"}:
                return "Polska"
            if candidate_norm in {"lt", "litwa", "lithuania"}:
                return "Litwa"
            return candidate

    return ""


def parse_vat_number(body: str) -> str:
    m_vat = re.search(r"Nr\s+VAT\s*:\s*([A-Z]{2}\d{6,})", body, flags=re.IGNORECASE)
    if m_vat:
        return m_vat.group(1).strip()
    return ""


def _drop_columns(ws, names: List[str]) -> None:
    indices: List[int] = []
    for name in names:
        idx = _find_column_index(ws, name)
        if idx is not None:
            indices.append(idx + 1)
    for idx in sorted(indices, reverse=True):
        ws.delete_cols(idx)


def _reorder_sheet(ws, headers: List[str]) -> None:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        ws.append(headers)
        return

    current_headers = [str(h or "").strip().lower() for h in rows[0]]
    index_map = {name: idx for idx, name in enumerate(current_headers) if name}
    new_rows = [headers]
    for row in rows[1:]:
        new_row = []
        for header in headers:
            idx = index_map.get(header.lower())
            new_row.append(row[idx] if idx is not None and idx < len(row) else None)
        new_rows.append(new_row)

    ws.delete_rows(1, ws.max_row)
    for row in new_rows:
        ws.append(row)


def _ensure_headers(ws, headers: List[str]) -> None:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row or all(value is None for value in header_row):
        ws.append(headers)
        return
    for idx, name in enumerate(headers, start=1):
        ws.cell(row=1, column=idx).value = name


def _shipping_international(country: str) -> str:
    if not country:
        return ""
    normalized = normalize_text(country).lower()
    if normalized in {"polska", "pl"}:
        return "nie"
    return "tak"


def _shipping_flag_from_country(country: str) -> str:
    if not country:
        return ""
    normalized = normalize_text(country).lower()
    if normalized in {"polska", "pl"}:
        return "nie"
    return "tak"


def parse_sale(msg: email.message.Message, body: str, subject: str, uid: str, translator: GoogleTranslator) -> Optional[Record]:
    tx = re.search(r"Numer\s+transakcji\s*[:#]?\s*#?\s*([0-9]{6,})", body, flags=re.IGNORECASE)
    tx_id = tx.group(1) if tx else uid

    title_match = re.search(
        r"Kupuj[aą]cy\s+potwierdzi[łl]\s+otrzymanie\s+(.+?)(?:\.|\n|$)",
        body,
        flags=re.IGNORECASE,
    )
    title_original = title_match.group(1).strip(" \"'“”") if title_match else ""
    if not title_original:
        alt = re.search(r"sprzedaz\s+zostala\s+zakonczona\.?\s*(.+)", normalize_text(body.lower()), flags=re.IGNORECASE)
        title_original = alt.group(1).strip() if alt else ""

    amount = ""
    m_wallet = re.search(
        r"Przelano\s+do\s+twojego\s+Portfela\s+Vinted\s*([0-9\s.,]+\s*z[łl])",
        body,
        flags=re.IGNORECASE,
    )
    if m_wallet:
        amount = m_wallet.group(1).strip()
    else:
        m_item = re.search(r"Otrzymano\s+za\s+przedmiot\s*[:]?\s*([0-9\s.,]+\s*z[łl])", body, flags=re.IGNORECASE)
        amount = m_item.group(1).strip() if m_item else ""

    title_pl = translate_to_polish(translator, title_original)
    country = parse_country(body)
    shipping_international = ""
    vat_number = parse_vat_number(body)

    return Record(
        kind="sprzedaz",
        tx_id=tx_id,
        email_uid=uid,
        subject=subject,
        title_original=title_original,
        title_pl=title_pl,
        amount_total=amount,
        date_text=parse_date_from_email(msg, body),
        country=country,
        shipping_international=shipping_international,
        vat_number=vat_number,
        transaction_number=tx_id,
    )


def parse_purchase(msg: email.message.Message, body: str, subject: str, uid: str, translator: GoogleTranslator) -> Optional[Record]:
    tx = re.search(r"Numer\s+transakcji\s*[:#]?\s*#?\s*([0-9]{6,})", body, flags=re.IGNORECASE)
    tx_id = tx.group(1) if tx else uid

    title_original = ""
    m_subject = re.search(r"Twoje\s+potwierdzenie\s+zakupu\s*:?\s*[\"“”']?(.+?)[\"“”']?$", subject, flags=re.IGNORECASE)
    if m_subject:
        title_original = m_subject.group(1).strip()

    if not title_original:
        m_order = re.search(r"Zam[oó]wienie\s+(.+?)\n", body, flags=re.IGNORECASE)
        if m_order:
            title_original = m_order.group(1).strip(" \"'“”")

    total = ""
    m_total = re.search(r"Zap[łl]acono\s*([0-9\s.,]+\s*z[łl])", body, flags=re.IGNORECASE)
    if m_total:
        total = m_total.group(1).strip()
    else:
        price = _extract_amount(body, r"Przedmiot\s*([0-9\s.,]+\s*z[łl])")
        shipping = _extract_amount(body, r"Wysy[łl]ka\s*([0-9\s.,]+\s*z[łl])")
        protection = _extract_amount(body, r"Op[łl]ata\s+za\s+Ochron[ęe]\s+Kupuj[ąa]cych\s*([0-9\s.,]+\s*z[łl])")
        total = sum_currency_strings([price, shipping, protection])

    title_pl = translate_to_polish(translator, title_original)
    country = parse_country(body)

    return Record(
        kind="zakup",
        tx_id=tx_id,
        email_uid=uid,
        subject=subject,
        title_original=title_original,
        title_pl=title_pl,
        amount_total=total,
        date_text=parse_date_from_email(msg, body),
        country=country,
        shipping_international="",
        vat_number="",
        transaction_number=tx_id,
    )


def parse_service(msg: email.message.Message, body: str, subject: str, uid: str) -> Optional[Record]:
    title = ""
    m_title = re.search(r"Usluga\s*([\w\s\-]+)", normalize_text(body), flags=re.IGNORECASE)
    if m_title:
        title = m_title.group(0).strip()
    if not title:
        title = subject.strip()

    amount = ""
    m_total = re.search(r"Kwota\s+ca[lł]kowita\s*([0-9\s.,]+\s*z[łl])", body, flags=re.IGNORECASE)
    if m_total:
        amount = m_total.group(1).strip()

    return Record(
        kind="usluga",
        tx_id=uid,
        email_uid=uid,
        subject=subject,
        title_original=title,
        title_pl="",
        amount_total=amount,
        date_text=parse_date_from_email(msg, body),
        country="",
        shipping_international="",
        vat_number="",
        transaction_number=uid,
    )


def _extract_amount(text: str, pattern: str) -> str:
    m = re.search(pattern, text, flags=re.IGNORECASE)
    return m.group(1).strip() if m else ""


def parse_money_to_float(value: str) -> Optional[float]:
    if not value:
        return None
    cleaned = re.sub(r"[^0-9,.-]", "", value).replace(" ", "")
    cleaned = cleaned.replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return None


def _to_excel_amount(value) -> Optional[float]:
    if isinstance(value, (int, float)):
        return float(value)
    return parse_money_to_float(str(value or ""))


def _to_excel_abs_amount(value) -> Optional[float]:
    num = _to_excel_amount(value)
    if num is None:
        return None
    return abs(num)


def _sanitize_title(value: str) -> str:
    text = _normalize_scraped_value(value)
    text = re.sub(r"^[\s,;:.\-–—'\"„”`]+", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _normalize_title_key(value: str) -> str:
    text = _sanitize_title(value).lower()
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _compact_title_key(value: str) -> str:
    normalized = _normalize_title_key(value)
    return re.sub(r"[\W_]+", "", normalized, flags=re.UNICODE)


def _titles_look_same(left: str, right: str) -> bool:
    a = _compact_title_key(left)
    b = _compact_title_key(right)
    if not a or not b:
        return False
    if a == b:
        return True
    if (a in b or b in a) and min(len(a), len(b)) >= 10:
        return True
    return False


def _is_bundle_title(value: str) -> bool:
    title = _normalize_title_key(value)
    return "zestaw" in title or "bundle" in title


def _normalize_amount_column(ws, header_name: str, force_abs: bool = False) -> None:
    idx = _find_column_index(ws, header_name)
    if idx is None:
        return
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or idx >= len(row):
            continue
        raw = row[idx]
        value = _to_excel_amount(raw)
        if value is None:
            continue
        if force_abs:
            value = abs(value)
        ws.cell(row=row_idx, column=idx + 1, value=value)


def _extract_date_only(value: str) -> str:
    text = str(value or "").strip()
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", text)
    if not m:
        return text
    dd = m.group(1).zfill(2)
    mm = m.group(2).zfill(2)
    yyyy = m.group(3)
    return f"{dd}.{mm}.{yyyy}"


def _date_sort_key(value: str) -> Tuple[int, int, int]:
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", str(value or ""))
    if not m:
        return (0, 0, 0)
    return (int(m.group(3)), int(m.group(2)), int(m.group(1)))


def _dedupe_and_sort_sheet(
    ws,
    title_header: str,
    amount_header: str,
    extra_headers: Optional[List[str]] = None,
    force_abs_amount: bool = False,
    merge_bundle_variants: bool = False,
    merge_similar_titles: bool = False,
) -> None:
    idx_date = _find_column_index(ws, "data")
    idx_title = _find_column_index(ws, title_header)
    idx_amount = _find_column_index(ws, amount_header)
    if idx_date is None or idx_title is None or idx_amount is None:
        return

    extra_headers = extra_headers or []
    extra_indices = [idx for idx in (_find_column_index(ws, h) for h in extra_headers) if idx is not None]
    extra_weights: Dict[int, int] = {}
    for header in extra_headers:
        idx = _find_column_index(ws, header)
        if idx is None:
            continue
        normalized = str(header or "").strip().lower()
        # Transaction number is the strongest identifier, prefer rows that contain it.
        extra_weights[idx] = 10 if normalized == "numer_transakcji" else 1

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        row_values = list(row)
        while len(row_values) <= max(idx_date, idx_title, idx_amount, *(extra_indices or [0])):
            row_values.append("")

        row_values[idx_date] = _extract_date_only(row_values[idx_date])
        amount_value = _to_excel_amount(row_values[idx_amount])
        if amount_value is not None and force_abs_amount:
            amount_value = abs(amount_value)
        row_values[idx_amount] = amount_value
        rows.append(row_values)

    bundle_date_amount_keys = set()
    if merge_bundle_variants:
        for row in rows:
            date_key = _extract_date_only(row[idx_date]).lower()
            amount_key = _amount_abs_key(str(row[idx_amount]))
            if _is_bundle_title(str(row[idx_title] or "")):
                bundle_date_amount_keys.add((date_key, amount_key))

    best_by_key: Dict[Tuple[str, str, str], List] = {}
    for row in rows:
        date_key = _extract_date_only(row[idx_date]).lower()
        title_key = _normalize_title_key(str(row[idx_title] or ""))
        amount_key = _amount_abs_key(str(row[idx_amount]))
        if merge_bundle_variants and (date_key, amount_key) in bundle_date_amount_keys:
            title_key = "__bundle__"
        key = (date_key, title_key, amount_key)
        if merge_similar_titles:
            for existing_key in list(best_by_key.keys()):
                if existing_key[0] != date_key or existing_key[2] != amount_key:
                    continue
                if _titles_look_same(title_key, existing_key[1]):
                    key = existing_key
                    break
        current = best_by_key.get(key)
        if current is None:
            best_by_key[key] = row
            continue
        current_score = sum(extra_weights.get(idx, 1) for idx in extra_indices if _normalize_scraped_value(current[idx]))
        new_score = sum(extra_weights.get(idx, 1) for idx in extra_indices if _normalize_scraped_value(row[idx]))
        if new_score > current_score:
            best_by_key[key] = row

    unique_rows = list(best_by_key.values())
    unique_rows.sort(key=lambda r: _date_sort_key(r[idx_date]))

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for row in unique_rows:
        ws.append(row)


def sum_currency_strings(values: List[str]) -> str:
    total = 0.0
    found = False
    for v in values:
        num = parse_money_to_float(v)
        if num is not None:
            total += num
            found = True
    if not found:
        return ""
    return f"{total:.2f}".replace(".", ",") + " zl"


def translate_to_polish(translator: GoogleTranslator, value: str) -> str:
    if not value:
        return ""
    try:
        translated = translator.translate(value)
        return translated.strip()
    except Exception:
        return value


def load_existing_ids(path: str) -> Tuple[set, set, set]:
    sale_ids = set()
    purchase_ids = set()
    service_ids = set()
    if not os.path.exists(path):
        return sale_ids, purchase_ids, service_ids

    wb = load_workbook(path)
    if "_Meta" in wb.sheetnames:
        ws = wb["_Meta"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            kind = str(row[0] or "").strip()
            key = str(row[1] or "").strip()
            if not kind or not key:
                continue
            if kind == "sprzedaz":
                sale_ids.add(key)
            elif kind == "zakup":
                purchase_ids.add(key)
            elif kind == "usluga":
                service_ids.add(key)

    if not sale_ids and "Sprzedaze" in wb.sheetnames:
        ws = wb["Sprzedaze"]
        key_idx = _find_column_index(ws, "email_uid")
        if key_idx is None:
            key_idx = _find_column_index(ws, "tx_id")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and key_idx is not None and row[key_idx]:
                sale_ids.add(str(row[key_idx]).strip())

    if not purchase_ids and "Zakupy" in wb.sheetnames:
        ws = wb["Zakupy"]
        key_idx = _find_column_index(ws, "email_uid")
        if key_idx is None:
            key_idx = _find_column_index(ws, "tx_id")
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and key_idx is not None and row[key_idx]:
                purchase_ids.add(str(row[key_idx]).strip())

    return sale_ids, purchase_ids, service_ids


def ensure_workbook(path: str) -> Workbook:
    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()

    purchase_headers = ["data", "tytul_oryginal", "tytul_pl", "kwota_lacznie"]
    if "Zakupy" not in wb.sheetnames:
        ws = wb.create_sheet("Zakupy")
        ws.append(purchase_headers)
    else:
        ws = wb["Zakupy"]
        _drop_columns(ws, ["temat", "email_uid", "tx_id", "kraj"])
        _reorder_sheet(ws, purchase_headers)

    sales_headers = [
        "data",
        "tytul_oryginal",
        "kwota",
        "numer_transakcji",
        "kraj_kupujacego",
        "wysylka_zagraniczna",
    ]
    if "Sprzedaze" not in wb.sheetnames:
        ws = wb.create_sheet("Sprzedaze")
        ws.append(sales_headers)
    else:
        ws = wb["Sprzedaze"]
        _drop_columns(ws, ["tytul_pl", "temat", "email_uid", "tx_id", "kraj", "nr_vat"])
        _reorder_sheet(ws, sales_headers)

    if "Uslugi elektroniczne" not in wb.sheetnames:
        ws = wb.create_sheet("Uslugi elektroniczne")
        ws.append(["data", "usluga", "kwota"])
    else:
        ws = wb["Uslugi elektroniczne"]
        _ensure_headers(ws, ["data", "usluga", "kwota"])

    if "Zwroty" not in wb.sheetnames:
        ws = wb.create_sheet("Zwroty")
        ws.append(["data", "tytul", "kwota"])
    else:
        ws = wb["Zwroty"]
        _ensure_headers(ws, ["data", "tytul", "kwota"])

    if "Podsumowanie" not in wb.sheetnames:
        ws = wb.create_sheet("Podsumowanie")
        ws.append(["pozycja", "suma"])

    if "_Meta" not in wb.sheetnames:
        ws = wb.create_sheet("_Meta")
        ws.append(["kind", "key"])
        ws.sheet_state = "hidden"

    _ensure_sheet_order(wb, ["Zakupy", "Sprzedaze", "Uslugi elektroniczne", "Zwroty", "Podsumowanie", "_Meta"])

    _normalize_amount_column(wb["Zakupy"], "kwota_lacznie", force_abs=True)
    _normalize_amount_column(wb["Sprzedaze"], "kwota")
    _normalize_amount_column(wb["Uslugi elektroniczne"], "kwota", force_abs=True)
    _normalize_amount_column(wb["Zwroty"], "kwota", force_abs=True)

    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 2:
        del wb["Sheet"]

    return wb


def append_record(ws, record: Record) -> None:
    if ws.title == "Sprzedaze":
        ws.append(
            [
                record.date_text,
                _sanitize_title(record.title_original),
                _to_excel_amount(record.amount_total),
                _normalize_scraped_value(record.transaction_number),
                "",
                record.shipping_international,
            ]
        )
        return

    if ws.title == "Uslugi elektroniczne":
        ws.append([record.date_text, _sanitize_title(record.title_original), _to_excel_abs_amount(record.amount_total)])
        return

    ws.append(
        [
            record.date_text,
            _sanitize_title(record.title_original),
            _sanitize_title(record.title_pl),
            _to_excel_abs_amount(record.amount_total),
        ]
    )


def record_key(record: Record) -> str:
    return record.tx_id or record.email_uid


def _find_column_index(ws, header_name: str) -> Optional[int]:
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        return None
    for idx, value in enumerate(header_row):
        if str(value or "").strip().lower() == header_name.lower():
            return idx
    return None


def _ensure_sheet_order(wb: Workbook, names: List[str]) -> None:
    existing = [name for name in names if name in wb.sheetnames]
    sheets = wb._sheets
    ordered = []
    for name in existing:
        for sh in sheets:
            if sh.title == name:
                ordered.append(sh)
                break
    for sh in sheets:
        if sh.title not in existing:
            ordered.append(sh)
    wb._sheets = ordered


def _autosize_columns(ws) -> None:
    lengths: Dict[int, int] = {}
    for row in ws.iter_rows(values_only=True):
        for idx, value in enumerate(row):
            text = str(value) if value is not None else ""
            lengths[idx] = max(lengths.get(idx, 0), len(text))
    for idx, length in lengths.items():
        col_letter = chr(ord("A") + idx)
        ws.column_dimensions[col_letter].width = min(max(length + 2, 10), 60)


def _html_escape(value: str) -> str:
        return (
                value.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace('"', "&quot;")
        )


def _sheet_to_html_table(ws, title: str) -> str:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
                return f"<h2>{_html_escape(title)}</h2><p>Brak danych.</p>"

        headers = rows[0]
        body_rows = rows[1:]

        header_cells = "".join(f"<th>{_html_escape(str(h or ''))}</th>" for h in headers)
        body_cells = []
        for row in body_rows:
                cells = "".join(f"<td>{_html_escape(str(c or ''))}</td>" for c in row)
                body_cells.append(f"<tr>{cells}</tr>")

        table_html = "\n".join(body_cells)
        return (
                f"<h2>{_html_escape(title)}</h2>"
                f"<table><thead><tr>{header_cells}</tr></thead>"
                f"<tbody>{table_html}</tbody></table>"
        )


def write_html_report(path: str, wb: Workbook) -> None:
    sales = wb["Sprzedaze"] if "Sprzedaze" in wb.sheetnames else None
    purchases = wb["Zakupy"] if "Zakupy" in wb.sheetnames else None
    services = wb["Uslugi elektroniczne"] if "Uslugi elektroniczne" in wb.sheetnames else None
    refunds = wb["Zwroty"] if "Zwroty" in wb.sheetnames else None
    summary = wb["Podsumowanie"] if "Podsumowanie" in wb.sheetnames else None

    sections = []
    if summary:
        sections.append(_sheet_to_html_table(summary, "Podsumowanie"))
    if purchases:
        sections.append(_sheet_to_html_table(purchases, "Zakupy"))
    if sales:
        sections.append(_sheet_to_html_table(sales, "Sprzedaze"))
    if services:
        sections.append(_sheet_to_html_table(services, "Uslugi elektroniczne"))
    if refunds:
        sections.append(_sheet_to_html_table(refunds, "Zwroty"))

        html = f"""<!doctype html>
<html lang="pl">
<head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <title>Vinted - Podsumowanie</title>
    <style>
        :root {{
            color-scheme: light;
            font-family: Arial, sans-serif;
            background: #f6f7f9;
            color: #1b1f24;
        }}
        body {{
            margin: 24px;
        }}
        h1 {{
            margin-bottom: 16px;
        }}
        h2 {{
            margin: 24px 0 8px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            background: #ffffff;
            box-shadow: 0 2px 8px rgba(0,0,0,0.08);
        }}
        th, td {{
            padding: 8px 10px;
            border-bottom: 1px solid #e2e5ea;
            text-align: left;
            vertical-align: top;
            font-size: 14px;
        }}
        th {{
            background: #eef1f6;
            position: sticky;
            top: 0;
            z-index: 1;
        }}
        tbody tr:nth-child(even) {{
            background: #f9fafb;
        }}
    </style>
</head>
<body>
    <h1>Vinted - Podsumowanie</h1>
    {"".join(sections)}
</body>
</html>"""

        with open(path, "w", encoding="utf-8") as handle:
                handle.write(html)


def _safe_imap_logout(mail) -> None:
    if mail is None:
        return
    try:
        state = getattr(mail, "state", "")
        if state == "LOGOUT":
            return
        mail.logout()
    except (OSError, TimeoutError, ssl.SSLError, imaplib.IMAP4.abort, imaplib.IMAP4.error):
        pass


def connect_imap(email_addr: str, password: str, imap_server: str, imap_port: int):
    retries = 3
    for attempt in range(1, retries + 1):
        try:
            mail = imaplib.IMAP4_SSL(imap_server, imap_port, timeout=60)
            try:
                mail.socket.settimeout(60)
            except Exception:
                pass
            mail.login(email_addr, password)
            mail.select("INBOX")
            return mail
        except (imaplib.IMAP4.abort, OSError, TimeoutError, ssl.SSLError) as exc:
            if attempt >= retries:
                raise
            print(f"Blad polaczenia IMAP ({exc}). Ponawiam {attempt}/{retries}...")
            time.sleep(2 * attempt)


def _fetch_imap_message(mail, uid: str, retries: int = 3):
    for attempt in range(1, retries + 1):
        try:
            status, msg_data = mail.fetch(uid, "(RFC822)")
            if status != "OK" or not msg_data:
                raise imaplib.IMAP4.error(f"Nieprawidlowy status fetch: {status}")
            return msg_data
        except (imaplib.IMAP4.abort, imaplib.IMAP4.error, OSError, TimeoutError, ssl.SSLError) as exc:
            if attempt >= retries:
                raise
            print(f"Blad pobierania wiadomosci {uid} ({exc}). Ponawiam {attempt}/{retries}...")
            try:
                mail.noop()
            except Exception:
                pass
            time.sleep(2 * attempt)


def filter_by_start_date(msg: email.message.Message, start_date: datetime) -> bool:
    date_hdr = msg.get("Date")
    if not date_hdr:
        return False
    try:
        dt = parsedate_to_datetime(date_hdr)
    except Exception:
        return False
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt >= start_date.astimezone(dt.tzinfo)


def _format_imap_date(value: datetime) -> str:
    return value.strftime("%d-%b-%Y")


def _month_range_for(ref: datetime) -> Tuple[datetime, datetime]:
    if ref.tzinfo is None:
        ref = ref.replace(tzinfo=timezone.utc)
    start = ref.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    if start.month == 12:
        end = start.replace(year=start.year + 1, month=1)
    else:
        end = start.replace(month=start.month + 1)
    return start, end


def _is_ascii(value: str) -> bool:
    try:
        value.encode("ascii")
        return True
    except UnicodeEncodeError:
        return False


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*]', "_", value)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned or "vinted_podsumowanie"


def _apply_template(template: str, month_str: str, label: str) -> str:
    return template.replace("{month}", month_str).replace("{period}", label)


def _build_dynamic_filename(prefix: str, month_str: str, ext: str) -> str:
    base = _safe_filename(f"{prefix} {month_str}")
    return f"{base}.{ext}"


def _is_yes(value: str) -> bool:
    normalized = re.sub(r"[^a-zA-Z]+", "", value or "").lower()
    return normalized in {"t", "tak", "y", "yes"}


def _should_use_imap(history_choice: str, has_email_credentials: bool) -> bool:
    if not has_email_credentials:
        return False
    return not _is_yes(history_choice)


def _search_uids(mail, start_date: datetime, end_date: Optional[datetime], sender_filter: Optional[str]) -> List[bytes]:
    criteria = ["SINCE", _format_imap_date(start_date)]
    if end_date is not None:
        criteria.extend(["BEFORE", _format_imap_date(end_date)])
    charset = None
    if sender_filter:
        criteria.extend(["FROM", f'"{sender_filter}"'])
        if not _is_ascii(sender_filter):
            charset = "UTF-8"

    status, data = mail.search(charset, *criteria)
    if status != "OK":
        raise RuntimeError("Nie udalo sie pobrac listy wiadomosci.")
    return data[0].split() if data and data[0] else []


def extract_records(mail, start_date: datetime, end_date: Optional[datetime], sender_contains: str) -> List[Record]:
    sender_filters = []
    if sender_contains:
        sender_filters.append(sender_contains)
        sender_filters.append(normalize_text(sender_contains))
    sender_filters.append("vinted")

    seen = set()
    uids: List[bytes] = []
    for sender_filter in sender_filters:
        if not sender_filter:
            continue
        batch = _search_uids(mail, start_date, end_date, sender_filter)
        for uid in batch:
            if uid not in seen:
                seen.add(uid)
                uids.append(uid)

    if not uids:
        uids = _search_uids(mail, start_date, end_date, None)
        print("Brak wynikow po filtrze nadawcy, uzywam tylko filtra daty.")

    print(f"Znaleziono wiadomosci w skrzynce (po filtrze): {len(uids)}")
    translator = GoogleTranslator(source="auto", target="pl")
    records: List[Record] = []

    sender_filters_norm = []
    if sender_contains:
        sender_filters_norm.append(normalize_text(sender_contains.lower()))
    sender_filters_norm.append("vinted")

    for uid_bytes in uids:
        uid = uid_bytes.decode("utf-8", errors="ignore")
        try:
            msg_data = _fetch_imap_message(mail, uid)
        except Exception:
            continue
        if not msg_data or not msg_data[0]:
            continue
        raw_email = msg_data[0][1]
        msg = email.message_from_bytes(raw_email)

        if not filter_by_start_date(msg, start_date):
            continue
        if end_date is not None:
            date_hdr = msg.get("Date")
            if not date_hdr:
                continue
            try:
                dt = parsedate_to_datetime(date_hdr)
                if dt.tzinfo is None:
                    dt = dt.replace(tzinfo=timezone.utc)
                if dt >= end_date.astimezone(dt.tzinfo):
                    continue
            except Exception:
                continue

        subject = decode_mime_header(msg.get("Subject", ""))
        sender = decode_mime_header(msg.get("From", ""))

        sender_norm = normalize_text(sender.lower())
        if not any(token and token in sender_norm for token in sender_filters_norm):
            continue

        subject_norm = normalize_text(subject.lower())

        body = extract_plain_text(msg)
        if not body:
            continue

        parsed: Optional[Record] = None
        if SALE_SUBJECT in subject_norm:
            parsed = parse_sale(msg, body, subject, uid, translator)
        elif PURCHASE_SUBJECT in subject_norm:
            parsed = parse_purchase(msg, body, subject, uid, translator)
        elif SERVICE_SUBJECT in subject_norm:
            parsed = parse_service(msg, body, subject, uid)

        if parsed:
            records.append(parsed)

    print(f"Znaleziono pasujace Vinted: {len(records)}")
    return records


def main() -> None:
    load_dotenv()

    email_addr = os.getenv("WP_EMAIL", "").strip()
    password = os.getenv("WP_PASSWORD", "").strip()
    imap_server = os.getenv("IMAP_SERVER", "imap.wp.pl").strip()
    imap_port = int(os.getenv("IMAP_PORT", "993"))
    output = os.getenv("OUTPUT_XLSX", "vinted_podsumowanie.xlsx").strip()
    output_html = os.getenv("OUTPUT_HTML", "vinted_podsumowanie.html").strip()
    dynamic_filename = os.getenv("DYNAMIC_FILENAME", "0").strip() == "1"
    output_prefix = os.getenv("OUTPUT_PREFIX", "podsumowanie Kamochi").strip()
    open_report = os.getenv("OPEN_REPORT", "0").strip() == "1"
    sender_contains = os.getenv("VINTED_FROM_CONTAINS", "vinted").strip()
    start_date_str = os.getenv("START_DATE", "2026-03-16").strip()

    history_only = False
    if not email_addr or not password:
        history_only = True
        print("Brak danych IMAP, uruchamiam tryb tylko historia Vinted.")

    try:
        start_date = datetime.strptime(start_date_str, "%Y-%m-%d")
    except ValueError as exc:
        raise RuntimeError("START_DATE musi miec format YYYY-MM-DD") from exc

    start_date = start_date.replace(tzinfo=timezone.utc)
    base_start_date = start_date
    end_date: Optional[datetime] = None
    period_label = f"START_DATE {start_date_str}"

    print("Wybierz okres skanowania:")
    print("  1) biezacy miesiac")
    print("  2) poprzedni miesiac")
    print("  3) START_DATE z .env")
    print("  4) podaj miesiac (YYYY-MM)")
    choice = input("Wybor (ENTER=1): ").strip()

    if choice in {"", "1"}:
        now = datetime.now(timezone.utc)
        start_date, end_date = _month_range_for(now)
        period_label = f"Biezacy miesiac {start_date.strftime('%Y-%m')}"
    elif choice == "2":
        now = datetime.now(timezone.utc)
        current_start, _ = _month_range_for(now)
        prev_month = current_start.replace(day=1) - timedelta(days=1)
        start_date, end_date = _month_range_for(prev_month)
        period_label = f"Poprzedni miesiac {start_date.strftime('%Y-%m')}"
    elif choice == "3":
        period_label = f"START_DATE {start_date_str}"
    elif choice == "4":
        month_input = input("Podaj miesiac (YYYY-MM): ").strip()
        try:
            month_start = datetime.strptime(month_input, "%Y-%m").replace(tzinfo=timezone.utc)
        except ValueError as exc:
            raise RuntimeError("Miesiac musi miec format YYYY-MM") from exc
        start_date, end_date = _month_range_for(month_start)
        period_label = f"Miesiac {start_date.strftime('%Y-%m')}"
    else:
        print("Nieznany wybor, uzywam START_DATE z .env.")
        period_label = f"START_DATE {start_date_str}"

    if start_date < base_start_date:
        start_date = base_start_date
        if end_date is not None and start_date >= end_date:
            print("Wybrany okres jest przed START_DATE, brak danych w tym miesiacu.")
        period_label = f"{period_label} (od START_DATE {start_date_str})"

    month_str = start_date.strftime("%Y-%m")
    default_xlsx = "vinted_podsumowanie.xlsx"
    default_html = "vinted_podsumowanie.html"
    use_dynamic = dynamic_filename or output == default_xlsx
    use_dynamic_html = dynamic_filename or output_html == default_html

    if "{month}" in output or "{period}" in output:
        output = _apply_template(output, month_str, period_label)
    elif use_dynamic:
        output = _build_dynamic_filename(output_prefix, month_str, "xlsx")

    if "{month}" in output_html or "{period}" in output_html:
        output_html = _apply_template(output_html, month_str, period_label)
    elif use_dynamic_html:
        output_html = _build_dynamic_filename(output_prefix, month_str, "html")

    output = _safe_filename(output)
    output_html = _safe_filename(output_html)

    history_choice = "t" if (not email_addr or not password) else input(
        "Uzyc tylko historie Vinted (bez maili IMAP)? (t/n): "
    ).strip()
    force_history = _is_yes(history_choice)
    history_only = history_only or force_history

    print("Start programu Vinted -> Excel")
    if end_date is None:
        print(f"Okres: {period_label}")
    else:
        print(f"Okres: {period_label} (od {start_date.strftime('%Y-%m-%d')})")
    if history_only:
        print("Tryb: tylko historia Vinted (bez IMAP)")
    else:
        print(f"IMAP: {imap_server}:{imap_port}")
    print(f"Excel: {output}")
    print(f"HTML: {output_html}")

    sale_ids, purchase_ids, service_ids = load_existing_ids(output)
    wb = ensure_workbook(output)
    ws_sales = wb["Sprzedaze"]
    ws_purchases = wb["Zakupy"]
    ws_services = wb["Uslugi elektroniczne"]
    ws_refunds = wb["Zwroty"]
    ws_summary = wb["Podsumowanie"]
    ws_meta = wb["_Meta"]

    records: List[Record] = []
    if _should_use_imap(history_choice, bool(email_addr and password)):
        print("Lacze z IMAP...")
        mail = connect_imap(email_addr, password, imap_server, imap_port)
        print("Polaczono. Pobieram i analizuje wiadomosci...")
        try:
            records = extract_records(mail, start_date, end_date, sender_contains)
        finally:
            try:
                mail.close()
            except Exception:
                pass
            _safe_imap_logout(mail)

    history_records: List[Dict[str, str]] = []
    if not force_history:
        history_choice = input(
            "Czy dociagnac historie zakupow/sprzedazy z Vinted (wallet/history)? (t/n): "
        ).strip()
    if _is_yes(history_choice):
        try:
            history_records = _run_puppeteer_history_scraper("vinted_scrape.js", month_str)
            print(f"Pobrano wpisy z historii Vinted: {len(history_records)}")
        except Exception as exc:
            print(f"Scraper historii nie powiodl sie: {exc}")

    cancel_remaining = _build_cancel_remaining(list(records) + list(history_records))

    added_sales = 0
    added_purchases = 0
    added_services = 0
    services_tuple_set = _build_services_existing_set(ws_services)

    for rec in records:
        rec.date_text = _extract_date_only(rec.date_text)
        key = record_key(rec)
        if rec.kind == "sprzedaz":
            if key in sale_ids:
                continue
            append_record(ws_sales, rec)
            sale_ids.add(key)
            ws_meta.append(["sprzedaz", key])
            added_sales += 1
        elif rec.kind == "zakup":
            balance_key = (
                _normalize_title_key(_normalize_scraped_value(rec.title_original)),
                _amount_abs_key(_normalize_scraped_value(rec.amount_total)),
            )
            if cancel_remaining.get(balance_key, 0) > 0:
                cancel_remaining[balance_key] -= 1
                continue
            if key in purchase_ids:
                continue
            append_record(ws_purchases, rec)
            purchase_ids.add(key)
            ws_meta.append(["zakup", key])
            added_purchases += 1
        elif rec.kind == "usluga":
            if key in service_ids:
                continue
            service_row_key = (
                _extract_date_only(_normalize_scraped_value(rec.date_text)).lower(),
                _normalize_title_key(rec.title_original),
                _amount_abs_key(_normalize_scraped_value(rec.amount_total)),
            )
            if service_row_key in services_tuple_set:
                continue
            append_record(ws_services, rec)
            service_ids.add(key)
            services_tuple_set.add(service_row_key)
            ws_meta.append(["usluga", key])
            added_services += 1

    filled_tx = _fill_sales_transaction_numbers_from_records(ws_sales, records)
    if filled_tx:
        print(f"Uzupelniono numery transakcji z maili: {filled_tx}")

    hist_added_sales, hist_added_purchases, hist_added_services, hist_added_refunds = _append_history_records(
        ws_sales, ws_purchases, ws_services, ws_refunds, history_records
    )
    filled_tx_history = _fill_sales_transaction_numbers_from_history(ws_sales, history_records)
    if filled_tx_history:
        print(f"Uzupelniono numery transakcji z historii: {filled_tx_history}")
    filled_country_history = _fill_sales_countries_from_history(ws_sales, history_records)
    if filled_country_history:
        print(f"Uzupelniono kraje kupujacych z historii: {filled_country_history}")
    if hist_added_sales or hist_added_purchases or hist_added_services or hist_added_refunds:
        print(
            "Historia Vinted -> dodano sprzedaze: "
            f"{hist_added_sales}, zakupy: {hist_added_purchases}, uslugi: {hist_added_services}, zwroty: {hist_added_refunds}"
        )
    added_sales += hist_added_sales
    added_purchases += hist_added_purchases
    added_services += hist_added_services

    targets = _build_scrape_targets(ws_sales)
    if not targets:
        print("Brak sprzedazy do uzupelnienia przez scraper.")
    else:
        print(f"Sprzedaze wymagajace uzupelnienia: {len(targets)}")
        scrape_choice = input("Czy zeskrapowac kraje kupujacych z Vinted? (t/n): ").strip()
        if _is_yes(scrape_choice):
            try:
                orders = _run_puppeteer_scraper("vinted_scrape.js", targets)
                updated = _update_sales_from_scrape(ws_sales, orders)
                print(f"Uzupelniono kraje/wysylke dla {updated} sprzedazy.")
            except Exception as exc:
                print(f"Scraper nie powiodl sie: {exc}")

    _dedupe_and_sort_sheet(
        ws_purchases,
        "tytul_oryginal",
        "kwota_lacznie",
        force_abs_amount=True,
        merge_bundle_variants=True,
    )
    _dedupe_and_sort_sheet(
        ws_sales,
        "tytul_oryginal",
        "kwota",
        ["numer_transakcji", "kraj_kupujacego", "wysylka_zagraniczna"],
        merge_similar_titles=True,
    )
    _dedupe_and_sort_sheet(ws_services, "usluga", "kwota", force_abs_amount=True)
    _dedupe_and_sort_sheet(ws_refunds, "tytul", "kwota", force_abs_amount=True)

    _update_summary(ws_summary, ws_sales, ws_purchases, ws_services, ws_refunds)

    _autosize_columns(ws_purchases)
    _autosize_columns(ws_sales)
    _autosize_columns(ws_services)
    _autosize_columns(ws_refunds)
    _autosize_columns(ws_summary)
    wb.save(output)
    write_html_report(output_html, wb)
    print(f"Dodano sprzedaze: {added_sales}")
    print(f"Dodano zakupy: {added_purchases}")
    print(f"Dodano uslugi elektroniczne: {added_services}")
    print(f"Plik zapisany: {output}")
    print(f"Raport HTML: {output_html}")

    if open_report:
        report_path = os.path.abspath(output_html)
        webbrowser.open(f"file:///{report_path}")


def _normalize_title_for_match(value: str) -> str:
    normalized = normalize_text(value or "").lower()
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return normalized


def _normalize_amount_for_match(value: str) -> Optional[float]:
    return parse_money_to_float(str(value or ""))


def _build_order_index(orders: List[Dict[str, str]]):
    by_tx: Dict[str, Dict[str, str]] = {}
    by_key: Dict[Tuple[str, float], List[Dict[str, str]]] = {}
    by_title: Dict[str, List[Dict[str, str]]] = {}
    by_price: Dict[float, List[Dict[str, str]]] = {}
    for order in orders:
        tx = str(order.get("transaction_number") or "").strip()
        if tx:
            by_tx[tx] = order

        match_type = str(order.get("match_type") or "").strip()
        title_key = _normalize_title_for_match(order.get("title", ""))
        amount_key = _normalize_amount_for_match(order.get("price", ""))

        if match_type == "title_price" and title_key and amount_key is not None:
            key = (title_key, round(amount_key, 2))
            by_key.setdefault(key, []).append(order)

        if match_type == "title_only" and title_key:
            by_title.setdefault(title_key, []).append(order)

        if match_type in {"price_only", "bundle_price"} and amount_key is not None:
            by_price.setdefault(round(amount_key, 2), []).append(order)

    return by_tx, by_key, by_title, by_price


def _fill_sales_transaction_numbers_from_records(ws_sales, records: List[Record]) -> int:
    idx_date = _find_column_index(ws_sales, "data")
    idx_title = _find_column_index(ws_sales, "tytul_oryginal")
    idx_amount = _find_column_index(ws_sales, "kwota")
    idx_tx = _find_column_index(ws_sales, "numer_transakcji")
    if idx_date is None or idx_title is None or idx_amount is None or idx_tx is None:
        return 0

    by_key: Dict[Tuple[str, str, str], List[str]] = {}
    by_title_amount: Dict[Tuple[str, str], List[str]] = {}
    for rec in records:
        if rec.kind != "sprzedaz":
            continue
        tx = _normalize_scraped_value(rec.transaction_number)
        if not tx:
            continue
        date_key = _extract_date_only(_normalize_scraped_value(rec.date_text)).lower()
        title_key = _normalize_title_key(rec.title_original)
        amount_key = _amount_abs_key(_normalize_scraped_value(rec.amount_total))
        by_key.setdefault((date_key, title_key, amount_key), []).append(tx)
        by_title_amount.setdefault((title_key, amount_key), []).append(tx)

    updated = 0
    for row_idx, row in enumerate(ws_sales.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        tx_val = row[idx_tx] if idx_tx < len(row) else ""
        if _normalize_scraped_value(tx_val):
            continue

        date_val = row[idx_date] if idx_date < len(row) else ""
        title_val = row[idx_title] if idx_title < len(row) else ""
        amount_val = row[idx_amount] if idx_amount < len(row) else ""
        key = (
            _extract_date_only(_normalize_scraped_value(date_val)).lower(),
            _normalize_title_key(title_val),
            _amount_abs_key(_normalize_scraped_value(amount_val)),
        )
        candidates = by_key.get(key, [])
        if not candidates:
            fallback_key = (key[1], key[2])
            candidates = by_title_amount.get(fallback_key, [])
        if not candidates:
            continue

        ws_sales.cell(row=row_idx, column=idx_tx + 1).value = candidates.pop(0)
        updated += 1

    return updated


def _fill_sales_transaction_numbers_from_history(ws_sales, history_records: List[Dict[str, str]]) -> int:
    idx_date = _find_column_index(ws_sales, "data")
    idx_title = _find_column_index(ws_sales, "tytul_oryginal")
    idx_amount = _find_column_index(ws_sales, "kwota")
    idx_tx = _find_column_index(ws_sales, "numer_transakcji")
    if idx_date is None or idx_title is None or idx_amount is None or idx_tx is None:
        return 0

    by_key: Dict[Tuple[str, str, str], List[str]] = {}
    by_title_amount: Dict[Tuple[str, str], List[str]] = {}
    for item in history_records:
        kind = _normalize_scraped_value(item.get("kind", "")).lower()
        if kind != "sprzedaz":
            continue
        tx = _normalize_scraped_value(item.get("transaction_number", ""))
        if not tx:
            tx = _fallback_tx_from_order_url(item.get("order_url", ""))
        if not tx:
            continue
        date_key = _extract_date_only(_normalize_scraped_value(item.get("date_text", ""))).lower()
        title_key = _normalize_title_key(item.get("title", ""))
        amount_key = _amount_abs_key(_normalize_scraped_value(item.get("price", "")))
        by_key.setdefault((date_key, title_key, amount_key), []).append(tx)
        by_title_amount.setdefault((title_key, amount_key), []).append(tx)

    updated = 0
    for row_idx, row in enumerate(ws_sales.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        tx_val = row[idx_tx] if idx_tx < len(row) else ""
        if _normalize_scraped_value(tx_val):
            continue

        date_val = row[idx_date] if idx_date < len(row) else ""
        title_val = row[idx_title] if idx_title < len(row) else ""
        amount_val = row[idx_amount] if idx_amount < len(row) else ""
        key = (
            _extract_date_only(_normalize_scraped_value(date_val)).lower(),
            _normalize_title_key(title_val),
            _amount_abs_key(_normalize_scraped_value(amount_val)),
        )
        candidates = by_key.get(key, [])
        if not candidates:
            fallback_key = (key[1], key[2])
            candidates = by_title_amount.get(fallback_key, [])
        if not candidates:
            continue

        ws_sales.cell(row=row_idx, column=idx_tx + 1).value = candidates.pop(0)
        updated += 1

    return updated


def _fill_sales_countries_from_history(ws_sales, history_records: List[Dict[str, str]]) -> int:
    idx_date = _find_column_index(ws_sales, "data")
    idx_title = _find_column_index(ws_sales, "tytul_oryginal")
    idx_amount = _find_column_index(ws_sales, "kwota")
    idx_country = _find_column_index(ws_sales, "kraj_kupujacego")
    idx_ship = _find_column_index(ws_sales, "wysylka_zagraniczna")
    if idx_title is None or idx_amount is None or idx_country is None:
        return 0

    by_key: Dict[Tuple[str, str, str], List[str]] = {}
    by_title_amount: Dict[Tuple[str, str], List[str]] = {}
    for item in history_records:
        kind = _normalize_scraped_value(item.get("kind", "")).lower()
        if kind != "sprzedaz":
            continue
        country = _normalize_scraped_value(item.get("country", ""))
        if not country:
            continue
        date_key = _extract_date_only(_normalize_scraped_value(item.get("date_text", ""))).lower()
        title_key = _normalize_title_key(item.get("title", ""))
        amount_key = _amount_abs_key(_normalize_scraped_value(item.get("price", "")))
        by_key.setdefault((date_key, title_key, amount_key), []).append(country)
        by_title_amount.setdefault((title_key, amount_key), []).append(country)

    updated = 0
    for row_idx, row in enumerate(ws_sales.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        existing = row[idx_country] if idx_country < len(row) else ""
        if _normalize_scraped_value(existing):
            continue

        date_val = row[idx_date] if idx_date is not None and idx_date < len(row) else ""
        title_val = row[idx_title] if idx_title < len(row) else ""
        amount_val = row[idx_amount] if idx_amount < len(row) else ""
        key = (
            _extract_date_only(_normalize_scraped_value(date_val)).lower(),
            _normalize_title_key(title_val),
            _amount_abs_key(_normalize_scraped_value(amount_val)),
        )
        candidates = by_key.get(key, [])
        if not candidates:
            candidates = by_title_amount.get((key[1], key[2]), [])
        if not candidates:
            continue

        country = candidates.pop(0)
        ws_sales.cell(row=row_idx, column=idx_country + 1).value = country
        if idx_ship is not None:
            ws_sales.cell(row=row_idx, column=idx_ship + 1).value = _shipping_flag_from_country(country)
        updated += 1

    return updated


def _update_sales_from_scrape(ws_sales, orders: List[Dict[str, str]]) -> int:
    if not orders:
        return 0
    idx_title = _find_column_index(ws_sales, "tytul_oryginal")
    idx_amount = _find_column_index(ws_sales, "kwota")
    idx_tx = _find_column_index(ws_sales, "numer_transakcji")
    idx_country = _find_column_index(ws_sales, "kraj_kupujacego")
    idx_ship = _find_column_index(ws_sales, "wysylka_zagraniczna")
    if idx_title is None or idx_amount is None:
        return 0

    by_tx, by_key, by_title, by_price = _build_order_index(orders)
    uncertain_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    updated = 0
    for row_idx, row in enumerate(ws_sales.iter_rows(min_row=2, values_only=False), start=2):
        row_values = [cell.value for cell in row]
        title = row_values[idx_title] if idx_title < len(row_values) else ""
        amount = row_values[idx_amount] if idx_amount < len(row_values) else ""
        tx = row_values[idx_tx] if idx_tx is not None and idx_tx < len(row_values) else ""

        order = None
        if tx:
            order = by_tx.get(str(tx).strip())

        if order is None:
            key = (
                _normalize_title_for_match(str(title or "")),
                _normalize_amount_for_match(amount) or 0.0,
            )
            candidates = by_key.get(key, [])
            if candidates:
                order = candidates.pop(0)

        if order is None:
            title_key = _normalize_title_for_match(str(title or ""))
            if title_key:
                candidates = by_title.get(title_key, [])
                if candidates:
                    order = candidates.pop(0)

        if order is None:
            price_key = _normalize_amount_for_match(amount)
            if price_key is not None:
                candidates = by_price.get(round(price_key, 2), [])
                if candidates:
                    order = candidates.pop(0)

        if not order:
            continue

        country = str(order.get("country") or "").strip()
        order_tx = str(order.get("transaction_number") or "").strip()
        match_type = str(order.get("match_type") or "").strip()
        uncertain = match_type in {"price_only", "bundle_price", "title_only"}
        if idx_tx is not None and order_tx:
            cell = ws_sales.cell(row=row_idx, column=idx_tx + 1)
            if not str(cell.value or "").strip():
                cell.value = order_tx
        if idx_country is not None and country:
            cell = ws_sales.cell(row=row_idx, column=idx_country + 1)
            cell.value = country
            if uncertain:
                cell.fill = uncertain_fill
        if idx_ship is not None and country:
            cell = ws_sales.cell(row=row_idx, column=idx_ship + 1)
            cell.value = _shipping_flag_from_country(country)
            if uncertain:
                cell.fill = uncertain_fill
        updated += 1

    return updated


def _build_scrape_targets(ws_sales) -> List[Dict[str, str]]:
    idx_title = _find_column_index(ws_sales, "tytul_oryginal")
    idx_amount = _find_column_index(ws_sales, "kwota")
    idx_tx = _find_column_index(ws_sales, "numer_transakcji")
    idx_country = _find_column_index(ws_sales, "kraj_kupujacego")
    idx_ship = _find_column_index(ws_sales, "wysylka_zagraniczna")
    if idx_title is None or idx_amount is None:
        return []

    targets: List[Dict[str, str]] = []
    for row in ws_sales.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        title = row[idx_title] if idx_title < len(row) else ""
        amount = row[idx_amount] if idx_amount < len(row) else ""
        tx = row[idx_tx] if idx_tx is not None and idx_tx < len(row) else ""
        country = row[idx_country] if idx_country is not None and idx_country < len(row) else ""
        ship = row[idx_ship] if idx_ship is not None and idx_ship < len(row) else ""

        if str(country or "").strip() and str(ship or "").strip():
            continue

        targets.append(
            {
                "title": str(title or ""),
                "amount": str(amount or ""),
                "transaction_number": str(tx or ""),
            }
        )

    return targets


def _run_puppeteer_scraper(script_path: str, targets: List[Dict[str, str]]) -> List[Dict[str, str]]:
    env = os.environ.copy()
    if targets:
        env["VINTED_TARGETS_JSON"] = json.dumps(targets, ensure_ascii=True)
    return _run_node_scraper(script_path, env)


def _run_puppeteer_history_scraper(script_path: str, target_month: str) -> List[Dict[str, str]]:
    env = os.environ.copy()
    env["VINTED_SCRAPE_MODE"] = "history"
    env["VINTED_SCRAPE_TARGET_MONTH"] = target_month
    return _run_node_scraper(script_path, env)


def _run_node_scraper(script_path: str, env: Dict[str, str]) -> List[Dict[str, str]]:
    if not os.path.exists(script_path):
        raise RuntimeError("Nie znaleziono vinted_scrape.js. Utworz plik i zainstaluj puppeteer.")

    print("Uruchamiam scraper Puppeteer. Zaloguj sie do Vinted, a potem potwierdz ENTER w terminalu.")
    print("Jesli logowanie przez Google jest blokowane, uzyj loginu/hasla Vinted albo ustaw VINTED_PROFILE_DIR.")
    result = subprocess.run(
        ["node", script_path],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
        env=env,
    )

    if result.returncode != 0:
        error_text = (result.stderr or "").strip() or (result.stdout or "").strip()
        raise RuntimeError(f"Blad scrapera: {error_text}")

    output = result.stdout or ""
    output_lines = [line for line in output.splitlines() if line.strip()]
    if not output_lines:
        return []

    if len(output_lines) > 1:
        for line in output_lines[:-1]:
            print(f"[scraper] {line}")

    json_line = ""
    for line in reversed(output_lines):
        stripped = line.strip()
        if stripped.startswith("[") or stripped.startswith("{"):
            json_line = stripped
            break

    if not json_line:
        return []

    try:
        return json.loads(json_line)
    except json.JSONDecodeError as exc:
        raise RuntimeError("Nie udalo sie odczytac JSON ze scrapera.") from exc


def _normalize_scraped_value(value: str) -> str:
    return str(value or "").replace("\xa0", " ").strip()


def _build_sales_existing_sets(ws_sales):
    idx_date = _find_column_index(ws_sales, "data")
    idx_title = _find_column_index(ws_sales, "tytul_oryginal")
    idx_amount = _find_column_index(ws_sales, "kwota")
    idx_tx = _find_column_index(ws_sales, "numer_transakcji")
    tx_set = set()
    tuple_set = set()
    for row in ws_sales.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        date_val = _extract_date_only(
            _normalize_scraped_value(row[idx_date] if idx_date is not None and idx_date < len(row) else "")
        )
        title_val = _normalize_title_key(row[idx_title] if idx_title is not None and idx_title < len(row) else "")
        amount_val = row[idx_amount] if idx_amount is not None and idx_amount < len(row) else ""
        amount_key = _amount_abs_key(str(amount_val))
        tx_val = _normalize_scraped_value(row[idx_tx] if idx_tx is not None and idx_tx < len(row) else "")
        if tx_val:
            tx_set.add(tx_val)
        if date_val or title_val or amount_key:
            tuple_set.add((date_val.lower(), title_val.lower(), amount_key))
    return tx_set, tuple_set


def _build_purchase_existing_set(ws_purchases):
    idx_date = _find_column_index(ws_purchases, "data")
    idx_title = _find_column_index(ws_purchases, "tytul_oryginal")
    idx_amount = _find_column_index(ws_purchases, "kwota_lacznie")
    existing = set()
    for row in ws_purchases.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        date_val = _extract_date_only(
            _normalize_scraped_value(row[idx_date] if idx_date is not None and idx_date < len(row) else "")
        )
        title_val = _normalize_title_key(row[idx_title] if idx_title is not None and idx_title < len(row) else "")
        amount_val = row[idx_amount] if idx_amount is not None and idx_amount < len(row) else ""
        amount_key = _amount_abs_key(str(amount_val))
        if date_val or title_val or amount_key:
            existing.add((date_val.lower(), title_val.lower(), amount_key))
    return existing


def _build_services_existing_set(ws_services):
    idx_date = _find_column_index(ws_services, "data")
    idx_name = _find_column_index(ws_services, "usluga")
    idx_amount = _find_column_index(ws_services, "kwota")
    existing = set()
    for row in ws_services.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        date_val = _extract_date_only(
            _normalize_scraped_value(row[idx_date] if idx_date is not None and idx_date < len(row) else "")
        )
        name_val = _normalize_title_key(row[idx_name] if idx_name is not None and idx_name < len(row) else "")
        amount_val = row[idx_amount] if idx_amount is not None and idx_amount < len(row) else ""
        amount_key = _amount_abs_key(str(amount_val))
        if date_val or name_val or amount_key:
            existing.add((date_val.lower(), name_val.lower(), amount_key))
    return existing


def _build_refunds_existing_set(ws_refunds):
    idx_date = _find_column_index(ws_refunds, "data")
    idx_title = _find_column_index(ws_refunds, "tytul")
    idx_amount = _find_column_index(ws_refunds, "kwota")
    existing = set()
    for row in ws_refunds.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        date_val = _extract_date_only(
            _normalize_scraped_value(row[idx_date] if idx_date is not None and idx_date < len(row) else "")
        )
        title_val = _normalize_title_key(row[idx_title] if idx_title is not None and idx_title < len(row) else "")
        amount_val = row[idx_amount] if idx_amount is not None and idx_amount < len(row) else ""
        amount_key = _amount_abs_key(str(amount_val))
        if date_val or title_val or amount_key:
            existing.add((date_val.lower(), title_val.lower(), amount_key))
    return existing


def _amount_abs_key(value: str) -> str:
    num = parse_money_to_float(str(value or ""))
    if num is None:
        return _normalize_scraped_value(value).lower()
    return f"{abs(num):.2f}"


def _amount_abs_display(value: str) -> Optional[float]:
    num = parse_money_to_float(str(value or ""))
    if num is None:
        return _to_excel_amount(value)
    return abs(num)


def _fallback_tx_from_order_url(order_url: str) -> str:
    url = _normalize_scraped_value(order_url)
    if not url:
        return ""
    m_item = re.search(r"/items/(\d+)", url, flags=re.IGNORECASE)
    if m_item:
        return m_item.group(1)
    m_any = re.search(r"(\d{6,})", url)
    if m_any:
        return m_any.group(1)
    return ""


def _build_cancel_remaining(records: List[object]) -> Dict[Tuple[str, str], int]:
    purchase_counts: Dict[Tuple[str, str], int] = {}
    refund_counts: Dict[Tuple[str, str], int] = {}
    for item in records:
        kind_value = item.kind if hasattr(item, "kind") else item.get("kind", "")
        kind = _normalize_scraped_value(kind_value).lower()
        if kind not in {"zakup", "zwrot"}:
            continue

        title_value = (
            getattr(item, "title_original", "")
            or getattr(item, "title", "")
            or (item.get("title_original") if hasattr(item, "get") else "")
            or (item.get("title") if hasattr(item, "get") else "")
        )
        amount_value = (
            getattr(item, "amount_total", "")
            or getattr(item, "price", "")
            or (item.get("amount_total") if hasattr(item, "get") else "")
            or (item.get("price") if hasattr(item, "get") else "")
        )
        title_key = _normalize_title_key(_normalize_scraped_value(title_value))
        amount_key = _amount_abs_key(_normalize_scraped_value(amount_value))
        key = (title_key, amount_key)

        if kind == "zakup":
            purchase_counts[key] = purchase_counts.get(key, 0) + 1
        elif kind == "zwrot":
            refund_counts[key] = refund_counts.get(key, 0) + 1

    cancel_remaining: Dict[Tuple[str, str], int] = {}
    for key, p_count in purchase_counts.items():
        r_count = refund_counts.get(key, 0)
        cancel_remaining[key] = min(p_count, r_count)
    return cancel_remaining


def _append_history_records(
    ws_sales, ws_purchases, ws_services, ws_refunds, history_records: List[Dict[str, str]]
) -> Tuple[int, int, int, int]:
    if not history_records:
        return 0, 0, 0, 0

    sales_tx_set, sales_tuple_set = _build_sales_existing_sets(ws_sales)
    purchase_tuple_set = _build_purchase_existing_set(ws_purchases)
    services_tuple_set = _build_services_existing_set(ws_services)
    refunds_tuple_set = _build_refunds_existing_set(ws_refunds)
    added_sales = 0
    added_purchases = 0
    added_services = 0
    added_refunds = 0

    cancel_remaining = _build_cancel_remaining(history_records)

    for item in history_records:
        kind = _normalize_scraped_value(item.get("kind", ""))
        title = _sanitize_title(item.get("title", ""))
        amount = _normalize_scraped_value(item.get("price", ""))
        date_text = _extract_date_only(_normalize_scraped_value(item.get("date_text", "")))
        tx = _normalize_scraped_value(item.get("transaction_number", ""))
        if not tx:
            tx = _fallback_tx_from_order_url(item.get("order_url", ""))
        country = _normalize_scraped_value(item.get("country", ""))
        amount_key = _amount_abs_key(amount)

        title_key = _normalize_title_key(title)
        row_key = (date_text.lower(), title_key, amount_key)
        if kind == "sprzedaz":
            if tx and tx in sales_tx_set:
                continue
            if row_key in sales_tuple_set:
                continue
            ws_sales.append(
                [
                    date_text,
                    title,
                    _to_excel_amount(amount),
                    tx,
                    country,
                    _shipping_flag_from_country(country) if country else "",
                ]
            )
            if tx:
                sales_tx_set.add(tx)
            sales_tuple_set.add(row_key)
            added_sales += 1
        elif kind == "zakup":
            balance_key = (title_key, _amount_abs_key(amount))
            if cancel_remaining.get(balance_key, 0) > 0:
                cancel_remaining[balance_key] -= 1
                continue
            if row_key in purchase_tuple_set:
                continue
            ws_purchases.append([date_text, title, title, _to_excel_abs_amount(amount)])
            purchase_tuple_set.add(row_key)
            added_purchases += 1
        elif kind == "usluga":
            service_row_key = (date_text.lower(), title_key, amount_key)
            if service_row_key in services_tuple_set:
                continue
            ws_services.append([date_text, title, _to_excel_abs_amount(amount)])
            services_tuple_set.add(service_row_key)
            added_services += 1
        elif kind == "zwrot":
            balance_key = (title_key, _amount_abs_key(amount))
            if cancel_remaining.get(balance_key, 0) > 0:
                cancel_remaining[balance_key] -= 1
                continue
            amount_display = _to_excel_abs_amount(amount)
            refund_amount_key = _amount_abs_key(str(amount_display if amount_display is not None else ""))
            refund_row_key = (date_text.lower(), title_key, refund_amount_key)
            if refund_row_key in refunds_tuple_set:
                continue
            ws_refunds.append([date_text, title, amount_display])
            refunds_tuple_set.add(refund_row_key)
            added_refunds += 1

    return added_sales, added_purchases, added_services, added_refunds


def _column_letter_by_header(ws, header_name: str) -> Optional[str]:
    idx = _find_column_index(ws, header_name)
    if idx is None:
        return None
    return get_column_letter(idx + 1)


def _update_summary(ws_summary, ws_sales, ws_purchases, ws_services, ws_refunds) -> None:
    rows = list(ws_summary.iter_rows(values_only=True))
    if rows:
        ws_summary.delete_rows(1, ws_summary.max_row)
    ws_summary.append(["pozycja", "suma"])

    sales_amount_col = _column_letter_by_header(ws_sales, "kwota")
    sales_ship_col = _column_letter_by_header(ws_sales, "wysylka_zagraniczna")
    purchases_amount_col = _column_letter_by_header(ws_purchases, "kwota_lacznie")
    services_amount_col = _column_letter_by_header(ws_services, "kwota")
    refunds_amount_col = _column_letter_by_header(ws_refunds, "kwota")

    sales_formula = f"=SUM(Sprzedaze!{sales_amount_col}:{sales_amount_col})" if sales_amount_col else "=0"
    if sales_amount_col and sales_ship_col:
        sales_int_formula = (
            f'=SUMIFS(Sprzedaze!{sales_amount_col}:{sales_amount_col},'
            f'Sprzedaze!{sales_ship_col}:{sales_ship_col},"tak")'
        )
    else:
        sales_int_formula = "=0"
    purchases_formula = f"=SUM(Zakupy!{purchases_amount_col}:{purchases_amount_col})" if purchases_amount_col else "=0"
    services_formula = (
        f"=SUM('Uslugi elektroniczne'!{services_amount_col}:{services_amount_col})" if services_amount_col else "=0"
    )
    refunds_formula = f"=SUM(Zwroty!{refunds_amount_col}:{refunds_amount_col})" if refunds_amount_col else "=0"

    ws_summary.append(["Suma sprzedazy", sales_formula])
    ws_summary.append(["Suma sprzedazy zagraniczna", sales_int_formula])
    ws_summary.append(["Suma zakupow", purchases_formula])
    ws_summary.append(["Suma uslug elektronicznych", services_formula])
    ws_summary.append(["Suma zwrotow", refunds_formula])


if __name__ == "__main__":
    main()
